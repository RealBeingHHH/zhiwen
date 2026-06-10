"""
了了之触 · 文档阅读器 — PDF/Word/Excel/HTML/CSV/JSON/TXT
读取任何格式的文档，提取文本和数据。
"""

import json
import os
import re
import subprocess
import tempfile
import urllib.request
import urllib.parse
from pathlib import Path
from typing import Optional


def read_document(source: str, max_chars: int = 15000) -> dict:
    """
    读取任何格式的文档。
    source: URL 或本地文件路径
    
    返回: {"format": "pdf"|"docx"|"xlsx"|..., "text": "...", "tables": [...], "error": None}
    """
    # 判断是 URL 还是本地文件
    is_url = source.startswith("http://") or source.startswith("https://")
    
    if is_url:
        # 下载到临时文件
        try:
            req = urllib.request.Request(source, headers={
                "User-Agent": "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36"
            })
            with urllib.request.urlopen(req, timeout=30) as resp:
                content = resp.read()
            
            # 从Content-Type或URL判断格式
            content_type = resp.headers.get("Content-Type", "")
            suffix = _guess_format(source, content_type)
            
            # 保存到临时文件
            tmp_path = f"/tmp/liaoliao_doc_{os.getpid()}.{suffix}"
            with open(tmp_path, "wb") as f:
                f.write(content)
        except Exception as e:
            return {"format": "unknown", "text": "", "tables": [], "error": str(e)}
    else:
        tmp_path = source
        suffix = Path(source).suffix.lower().lstrip(".")
    
    # 按格式读取
    result = _read_by_format(tmp_path, suffix, max_chars)
    
    # 清理临时文件
    if is_url:
        try:
            os.remove(tmp_path)
        except:
            pass
    
    return result


def _guess_format(url: str, content_type: str) -> str:
    """猜测文档格式。"""
    url_lower = url.lower()
    
    if "pdf" in content_type or url_lower.endswith(".pdf"):
        return "pdf"
    if "epub" in content_type or url_lower.endswith(".epub"):
        return "epub"
    if "word" in content_type or "docx" in content_type or url_lower.endswith(".docx"):
        return "docx"
    if "excel" in content_type or "xlsx" in content_type or url_lower.endswith(".xlsx") or url_lower.endswith(".xls"):
        return "xlsx"
    if "csv" in content_type or url_lower.endswith(".csv"):
        return "csv"
    if "json" in content_type or url_lower.endswith(".json"):
        return "json"
    if "html" in content_type or url_lower.endswith(".html") or url_lower.endswith(".htm"):
        return "html"
    if url_lower.endswith(".txt") or url_lower.endswith(".md"):
        return "txt"
    
    # 最后看URL路径
    for fmt in ["pdf", "docx", "xlsx", "csv", "json", "html", "txt"]:
        if f".{fmt}" in url_lower:
            return fmt
    
    return "html"  # 默认按HTML处理


def _read_by_format(path: str, fmt: str, max_chars: int) -> dict:
    """按格式读取文件。"""
    if fmt == "pdf":
        return _read_pdf(path, max_chars)
    elif fmt == "epub":
        return _read_epub(path, max_chars)
    elif fmt in ("docx", "doc"):
        return _read_docx(path, max_chars)
    elif fmt in ("xlsx", "xls"):
        return _read_xlsx(path, max_chars)
    elif fmt == "csv":
        return _read_csv(path, max_chars)
    elif fmt == "json":
        return _read_json(path, max_chars)
    elif fmt in ("html", "htm"):
        return _read_html(path, max_chars)
    elif fmt == "txt":
        return _read_txt(path, max_chars)
    else:
        # 尝试所有格式
        for f in ["pdf", "docx", "xlsx", "csv", "json", "html", "txt"]:
            result = _read_by_format(path, f, max_chars)
            if result.get("text") and len(result["text"]) > 100:
                return result
        return {"format": fmt, "text": "", "tables": [], "error": f"不支持格式: {fmt}"}


def _read_pdf(path: str, max_chars: int) -> dict:
    """读取PDF。"""
    # 方案1: pymupdf
    try:
        import fitz
        doc = fitz.open(path)
        text_parts = []
        for page in doc:
            text = page.get_text()
            if text.strip():
                text_parts.append(text)
        doc.close()
        return {"format": "pdf", "text": "\n".join(text_parts)[:max_chars], "tables": [], "pages": len(text_parts)}
    except ImportError:
        pass

    # 方案2: pdftotext 命令行
    import subprocess as _sp
    try:
        r = _sp.run(['pdftotext', '-layout', path, '-', '-f', '1', '-l', '10'], capture_output=True, text=True, timeout=30)
        if r.returncode == 0 and r.stdout.strip():
            return {"format": "pdf", "text": r.stdout[:max_chars], "tables": [], "pages": "1-10"}
    except Exception:
        pass

    return {"format": "pdf", "text": "", "tables": [], "error": "无PDF阅读器(pymupdf/pdftotext)"}


def _read_epub(path: str, max_chars: int) -> dict:
    """读取EPUB电子书。解压ZIP→提取HTML文本。"""
    import zipfile, re as _re_epub
    try:
        with zipfile.ZipFile(path) as z:
            text_parts = []
            chapters = []
            for f in sorted(z.namelist()):
                if f.endswith(('.html', '.xhtml', '.htm')):
                    content = z.read(f).decode('utf-8', errors='ignore')
                    text = _re_epub.sub(r'<[^>]+>', ' ', content)
                    text = _re_epub.sub(r'\s+', ' ', text).strip()
                    if len(text) > 100:
                        text_parts.append(text)
                        # 提取章节标题
                        title_match = _re_epub.search(r'第[一二三四五六七八九十\d]+章\s*[^\s]{2,30}', text)
                        if title_match:
                            chapters.append(title_match.group(0))
            
            full_text = "\n\n---\n\n".join(text_parts)
            return {
                "format": "epub",
                "text": full_text[:max_chars],
                "tables": [],
                "chapters": chapters[:20],
                "sections": len(text_parts),
            }
    except Exception as e:
        return {"format": "epub", "text": "", "tables": [], "error": str(e)}


def _read_docx(path: str, max_chars: int) -> dict:
    """读取Word文档。"""
    try:
        from docx import Document
        doc = Document(path)
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        
        tables = []
        for t in doc.tables:
            rows = []
            for row in t.rows:
                rows.append([cell.text.strip() for cell in row.cells])
            tables.append(rows)
        
        return {"format": "docx", "text": text[:max_chars], "tables": tables[:5]}
    except ImportError:
        # 回退：用 python-docx 命令行
        try:
            result = subprocess.run(
                ["python3", "-c", f"from docx import Document; d=Document('{path}'); print('\\n'.join(p.text for p in d.paragraphs))"],
                capture_output=True, text=True, timeout=15
            )
            if result.returncode == 0:
                return {"format": "docx", "text": result.stdout[:max_chars], "tables": []}
        except:
            pass
        return {"format": "docx", "text": "", "tables": [], "error": "python-docx未安装"}
    except Exception as e:
        return {"format": "docx", "text": "", "tables": [], "error": str(e)}


def _read_xlsx(path: str, max_chars: int) -> dict:
    """读取Excel表格。"""
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=True)
        
        all_text = []
        tables = []
        
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = []
            for row in ws.iter_rows(values_only=True):
                row_data = [str(c) if c is not None else "" for c in row]
                if any(row_data):
                    rows.append(row_data)
                    all_text.append(" | ".join(row_data))
            
            if rows:
                tables.append({"sheet": sheet_name, "rows": rows[:50]})
        
        text = "\n".join(all_text[:200])
        return {"format": "xlsx", "text": text[:max_chars], "tables": tables[:3], "sheets": wb.sheetnames}
    except ImportError:
        return {"format": "xlsx", "text": "", "tables": [], "error": "openpyxl未安装"}
    except Exception as e:
        return {"format": "xlsx", "text": "", "tables": [], "error": str(e)}


def _read_csv(path: str, max_chars: int) -> dict:
    """读取CSV。"""
    try:
        with open(path, "r", encoding="utf-8-sig", errors="ignore") as f:
            content = f.read()
        return {"format": "csv", "text": content[:max_chars], "tables": []}
    except Exception as e:
        return {"format": "csv", "text": "", "tables": [], "error": str(e)}


def _read_json(path: str, max_chars: int) -> dict:
    """读取JSON。"""
    try:
        with open(path, "r") as f:
            data = json.load(f)
        text = json.dumps(data, ensure_ascii=False, indent=2)
        return {"format": "json", "text": text[:max_chars], "tables": []}
    except Exception as e:
        return {"format": "json", "text": "", "tables": [], "error": str(e)}


def _read_html(path: str, max_chars: int) -> dict:
    """读取HTML。"""
    try:
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            html = f.read()
        text = re.sub(r'<script[^>]*>.*?</script>', '', html, flags=re.DOTALL|re.IGNORECASE)
        text = re.sub(r'<style[^>]*>.*?</style>', '', text, flags=re.DOTALL|re.IGNORECASE)
        text = re.sub(r'<[^>]+>', '\n', text)
        text = re.sub(r'\n{3,}', '\n\n', text)
        text = re.sub(r'[ \t]+', ' ', text)
        return {"format": "html", "text": text.strip()[:max_chars], "tables": []}
    except Exception as e:
        return {"format": "html", "text": "", "tables": [], "error": str(e)}


def _read_txt(path: str, max_chars: int) -> dict:
    """读取文本文件。"""
    try:
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read()
        return {"format": "txt", "text": text[:max_chars], "tables": []}
    except Exception as e:
        return {"format": "txt", "text": "", "tables": [], "error": str(e)}


def extract_financial_data(text: str) -> dict:
    """从文本中提取财务数据。优先匹配五年对照表。"""
    results = {}

    # 先找 "FIVE-YEAR" 或 "五年" 标志，从那之后开始匹配
    start_idx = 0
    for marker in ["FIVE-YEAR", "五年主要財務數據", "五年主要财务数据"]:
        idx = text.find(marker)
        if idx > 0:
            start_idx = idx
            break

    # 在表格区域搜索
    table_text = text[start_idx:start_idx + 3000] if start_idx > 0 else text

    patterns = {
        "revenue": [
            r'Revenue\s+營業額\s+([\d,]+)\s+([\d,]+)',  # 取前两个数字(2025, 2024)
        ],
        "net_profit": [
            r'(?:Profit attributable|母公司擁有人\s*應佔溢利)\s+([\d,]+)\s+([\d,]+)',
        ],
        "gross_profit": [
            r'Gross profit\s+毛利\s+([\d,]+)\s+([\d,]+)',
        ],
        "gross_margin": [
            r'Gross profit margin.*?(\d+)\s+(\d+)',
        ],
        "net_margin": [
            r'Net profit margin.*?([\d.]+)\s+([\d.]+)',
        ],
    }

    for key, pats in patterns.items():
        for pat in pats:
            m = re.search(pat, table_text, re.IGNORECASE | re.DOTALL)
            if m:
                if len(m.groups()) >= 2:
                    results[key] = f"{m.group(1)} (2025) | {m.group(2)} (2024)"
                else:
                    results[key] = m.group(1)
                break

    # 如果表格区域没找到，在全文中找
    if not results:
        for pat in [r'(?:营收|营业收入|Revenue|营业额)[^\d]*(\d[\d,.]*\s*(?:亿|万))']:
            m = re.search(pat, text, re.IGNORECASE)
            if m:
                results["revenue"] = m.group(1)
                break

    return results
